"""
CHUN KING SharePoint to Google Drive Weekly Backup Pipeline
Orchestrates list exports (Excel + JSON), media harvesting,
throttling defense, incremental sync via manifest, and Google Drive upload.
"""

import os
import sys
import time
import json
import re
import datetime
from typing import Dict, Any, List, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from sharepoint_client import SharePointClient, decode_sharepoint_field, SharePointThrottlingException
from gdrive_client import GoogleDriveClient

# Reconfigure standard output for UTF-8 logging in CI/CD runners
sys.stdout.reconfigure(encoding="utf-8", line_buffering=True)

# Configuration & Constants
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_DIR = os.path.dirname(SCRIPT_DIR)
SITE_URL = os.environ.get("SP_SITE_URL", "https://k35n.sharepoint.com/sites/CHUNKING")

GUID_PROJECTS = "1958fe5e-336d-43a0-beb2-4da48df59f7b"
GUID_EXPENSES = "a09172c2-2be4-4b12-8dfd-418a6fbe0c6d"
GUID_INCOME = "527698fd-139d-4482-b819-b3f6da4e8794"
GUID_RECEIPTS_LIB = "441ac1bf-d867-4c5f-af6e-6939afcbb2b4"


def load_sharepoint_cookies() -> Dict[str, str]:
    """Loads SharePoint cookies from environment variable or local cookies.json."""
    env_cookies = os.environ.get("SP_COOKIES_JSON")
    if env_cookies:
        try:
            parsed = json.loads(env_cookies)
            if isinstance(parsed, list):
                return {c["name"]: c["value"] for c in parsed}
            return parsed
        except Exception as e:
            print(f"⚠️ Could not parse SP_COOKIES_JSON environment variable: {e}")

    local_cookie_path = os.path.join(PROJECT_DIR, "cookies.json")
    if os.path.exists(local_cookie_path):
        with open(local_cookie_path, "r", encoding="utf-8") as f:
            parsed = json.load(f)
            if isinstance(parsed, list):
                return {c["name"]: c["value"] for c in parsed}
            return parsed

    raise ValueError("No SharePoint cookies found in SP_COOKIES_JSON env var or local cookies.json!")


def clean_row_for_export(raw_item: Dict[str, Any]) -> Dict[str, Any]:
    """Cleans OData metadata and translates hex column keys to Traditional Chinese."""
    cleaned = {}
    for k, v in raw_item.items():
        if k in ("__metadata", "ContentTypeID", "FileSystemObjectType", "ServerRedirectedEmbedUri"):
            continue
        
        # Translate header
        cn_key = decode_sharepoint_field(k)
        
        # Format values
        if isinstance(v, dict):
            if "Url" in v:  # SharePoint URL field
                cleaned[cn_key] = v.get("Url", "")
            elif "results" in v:
                cleaned[cn_key] = json.dumps(v["results"], ensure_ascii=False)
            else:
                cleaned[cn_key] = str(v)
        else:
            cleaned[cn_key] = v
            
    return cleaned


def build_excel_sheet(workbook: openpyxl.Workbook, sheet_title: str, rows: List[Dict[str, Any]], primary_fields: List[str]):
    """Creates a beautifully styled Excel worksheet with headers and formatting."""
    ws = workbook.create_sheet(title=sheet_title)
    
    if not rows:
        ws.append(["無記錄"])
        return

    # Determine columns: primary fields first, then any remaining fields
    all_keys = set()
    for r in rows:
        all_keys.update(r.keys())
        
    ordered_headers = [f for f in primary_fields if f in all_keys]
    remaining_headers = sorted([f for f in all_keys if f not in ordered_headers])
    headers = ordered_headers + remaining_headers
    
    # Header styling
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border_thin = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    # Add data rows
    data_font = Font(name="Segoe UI", size=10)
    for row_idx, r in enumerate(rows, start=2):
        row_values = [r.get(h, "") for h in headers]
        ws.append(row_values)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = border_thin
            # Auto format currency if numeric
            val = cell.value
            if isinstance(val, (int, float)) and ("金額" in headers[col_idx-1] or "成本" in headers[col_idx-1] or "價錢" in headers[col_idx-1]):
                cell.number_format = '"HK$"#,##0.00'

    # Auto adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            v_str = str(cell.value or "")
            max_len = max(max_len, min(len(v_str), 40))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)


def run_backup_pipeline(lists_only: bool = False, max_media_files: Optional[int] = None):
    start_time = time.time()
    today_str = datetime.datetime.now().strftime("%Y-%m-%d")
    timestamp_full = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S HKT")

    print("==================================================================")
    print(f"🚀 CHUN KING SharePoint to Google Drive Weekly Backup Pipeline")
    print(f"📅 Timestamp: {timestamp_full}")
    print(f"🌐 SharePoint Site: {SITE_URL}")
    print("==================================================================")

    # 1. Initialize Clients
    cookies = load_sharepoint_cookies()
    sp = SharePointClient(
        site_url=SITE_URL,
        cookies=cookies,
        user_agent="NONISV|ChunKing|WeeklyBackup/1.0",
        pacing_delay=0.08,               # 80ms politeness delay between requests
        max_consecutive_throttles=5      # Circuit breaker threshold
    )

    gdrive_json = os.environ.get("GDRIVE_SERVICE_ACCOUNT_JSON")
    gdrive_folder_id = os.environ.get("GDRIVE_FOLDER_ID")
    local_service_account_file = os.path.join(SCRIPT_DIR, "service_account.json")
    
    oauth_json = os.environ.get("GDRIVE_OAUTH_TOKEN_JSON")
    local_oauth_file = os.path.join(SCRIPT_DIR, "oauth_token.json")

    gdrive = GoogleDriveClient(
        service_account_json_content=gdrive_json,
        service_account_file_path=local_service_account_file if os.path.exists(local_service_account_file) else None,
        oauth_json_content=oauth_json,
        oauth_file_path=local_oauth_file if os.path.exists(local_oauth_file) else None,
        root_folder_id=gdrive_folder_id
    )

    if gdrive.is_connected:
        print(f"🔗 Google Drive Connection: CONNECTED (Auth Mode: {gdrive.auth_mode})")
    else:
        print("⚠️ Google Drive Connection: LOCAL STAGING MODE (Files will be stored locally in backup/staging/)")

    staging_dir = os.path.join(SCRIPT_DIR, "staging", today_str)
    os.makedirs(staging_dir, exist_ok=True)
    temp_receipts_dir = os.path.join(staging_dir, "temp_downloads")
    os.makedirs(temp_receipts_dir, exist_ok=True)

    # 2. Fetch or initialize Manifest
    manifest = {}
    if gdrive.is_connected:
        manifest = gdrive.fetch_manifest()
        print(f"📦 Loaded remote manifest from Google Drive ({len(manifest.get('files', {}))} tracked files).")
    else:
        local_manifest_path = os.path.join(SCRIPT_DIR, "manifest.json")
        if os.path.exists(local_manifest_path):
            with open(local_manifest_path, "r", encoding="utf-8") as f:
                manifest = json.load(f)

    manifest_files = manifest.setdefault("files", {})

    # Target folders on Google Drive
    gdrive_weekly_folder_id = None
    gdrive_expense_folder_ids = {}
    gdrive_receipts_inc_id = None
    expense_subdirs = ["CHQ", "CK", "DBS", "WS"]

    if gdrive.is_connected:
        snapshots_root_id = gdrive.find_or_create_folder("Weekly_Snapshots")
        gdrive_weekly_folder_id = gdrive.find_or_create_folder(today_str, parent_id=snapshots_root_id)
        
        mirror_root_id = gdrive.find_or_create_folder("Receipts_Live_Mirror")
        for sub_name in expense_subdirs:
            gdrive_expense_folder_ids[sub_name] = gdrive.find_or_create_folder(sub_name, parent_id=mirror_root_id)
        gdrive_receipts_inc_id = gdrive.find_or_create_folder("Income", parent_id=mirror_root_id)

    # =========================================================================
    # 3. EXPORT SHAREPOINT LISTS (Full Snapshot)
    # =========================================================================
    print("\n------------------------------------------------------------------")
    print("📋 [Step 1/3] Extracting Structured Lists from SharePoint Online...")
    print("------------------------------------------------------------------")

    # 3.1 Projects List
    print("⏳ Pulling Projects (工程項目)...")
    raw_projects = sp.fetch_list_items(GUID_PROJECTS)
    clean_projects = [clean_row_for_export(it) for it in raw_projects]
    print(f"   ✅ Projects: {len(clean_projects)} records extracted.")

    # 3.2 Transactions / Expenses List
    print("⏳ Pulling Transactions (支出與支票)...")
    raw_expenses = sp.fetch_list_items(GUID_EXPENSES)
    clean_expenses = [clean_row_for_export(it) for it in raw_expenses]
    print(f"   ✅ Expenses: {len(clean_expenses)} records extracted.")

    # 3.3 Income List
    print("⏳ Pulling Income (收入記錄)...")
    raw_incomes = sp.fetch_list_items(GUID_INCOME)
    clean_incomes = [clean_row_for_export(it) for it in raw_incomes]
    print(f"   ✅ Income: {len(clean_incomes)} records extracted.")

    # 3.4 Generate Combined & Formatted Excel Workbook
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default blank sheet

    build_excel_sheet(
        wb, "工程項目", clean_projects,
        ["Id", "項目編號", "Title", "分類", "負責人", "合約價錢", "已收金額", "未收金額", "已付成本", "流入／出", "預計施工日期", "預計完工日期", "備註"]
    )
    build_excel_sheet(
        wb, "支出明細", clean_expenses,
        ["Id", "交易時間", "金額", "支票編號", "轉帳銀行", "所屬項目", "所屬項目 (ID)", "子項目", "收款人", "付款狀態", "收據連結", "RecordStatus", "IsIndependent", "BatchId"]
    )
    build_excel_sheet(
        wb, "收入記錄", clean_incomes,
        ["Id", "交易時間", "金額", "收項銀行", "所屬項目", "所屬項目 (ID)", "收據", "RecordStatus", "Remarks"]
    )

    master_excel_path = os.path.join(staging_dir, f"CHUN_KING_Financial_Master_{today_str}.xlsx")
    wb.save(master_excel_path)
    print(f"📊 Formatted Master Excel created: {os.path.basename(master_excel_path)}")

    # 3.5 Save Raw JSON data for lossless recovery
    raw_bundle_path = os.path.join(staging_dir, f"raw_data_{today_str}.json")
    with open(raw_bundle_path, "w", encoding="utf-8") as f:
        json.dump({
            "backup_date": today_str,
            "counts": {
                "projects": len(clean_projects),
                "expenses": len(clean_expenses),
                "incomes": len(clean_incomes)
            },
            "projects": clean_projects,
            "expenses": clean_expenses,
            "incomes": clean_incomes
        }, f, ensure_ascii=False, indent=2)

    # Upload List Snapshot to Google Drive
    if gdrive.is_connected:
        print("☁️ Uploading List Snapshot to Google Drive Weekly Folder...")
        gdrive.upload_file(master_excel_path, destination_folder_id=gdrive_weekly_folder_id, overwrite=True)
        gdrive.upload_file(raw_bundle_path, destination_folder_id=gdrive_weekly_folder_id, overwrite=True)
        print("   ✅ Lists safely archived on Google Drive.")

    # =========================================================================
    # 4. INCREMENTAL SYNC OF MEDIA & RECEIPT FILES (PDF / PNG / JPG)
    # =========================================================================
    print("\n------------------------------------------------------------------")
    print("📁 [Step 2/3] Incremental Synchronization of Receipts (PDF/Images)...")
    print("------------------------------------------------------------------")

    new_downloads_count = 0
    new_bytes_transferred = 0
    skipped_count = 0

    if lists_only:
        print("⏩ Skipping media synchronization (--lists-only flag is set).")
    else:
        # 4.1 Sync Expense Cheque Receipt Folders (CHQ, CK, DBS, WS)
        for folder_name in expense_subdirs:
            if max_media_files and new_downloads_count >= max_media_files:
                print(f"⏸️ Reached maximum media files limit ({max_media_files}). Stopping media sync.")
                break

            sp_folder_path = f"/sites/CHUNKING/DocLib/{folder_name}"
            target_gdrive_folder_id = gdrive_expense_folder_ids.get(folder_name)

            # Pre-scan Google Drive's folder in case previous run was cancelled before saving manifest
            existing_drive_files = {}
            if gdrive.is_connected and target_gdrive_folder_id:
                print(f"\n🔍 Pre-scanning existing files in Google Drive {folder_name} folder...")
                existing_drive_files = gdrive.list_files_in_folder(target_gdrive_folder_id)
                print(f"   Found {len(existing_drive_files)} files already present in Google Drive {folder_name} folder.")

            print(f"⏳ Scanning SharePoint {sp_folder_path}...")
            folder_files = sp.fetch_folder_files_metadata(sp_folder_path)
            print(f"   Found {len(folder_files)} total files in {folder_name} document library.")

            folder_new_count = 0
            folder_skipped_count = 0

            for file_meta in folder_files:
                if max_media_files and new_downloads_count >= max_media_files:
                    print(f"⏸️ Reached maximum media files limit ({max_media_files}). Stopping media sync.")
                    break

                name = file_meta.get("Name", "")
                rel_url = file_meta.get("ServerRelativeUrl", "")
                file_len = int(file_meta.get("Length", 0))

                # Check 1: Manifest check
                manifest_entry = manifest_files.get(rel_url)
                if manifest_entry and manifest_entry.get("size") == file_len:
                    skipped_count += 1
                    folder_skipped_count += 1
                    continue

                # Check 2: Direct Google Drive check (handles previously cancelled or interrupted runs)
                if name in existing_drive_files and existing_drive_files[name] == file_len:
                    manifest_files[rel_url] = {
                        "name": name,
                        "size": file_len,
                        "sha256": "pre-existing",
                        "category": folder_name,
                        "last_backed_up": today_str
                    }
                    skipped_count += 1
                    folder_skipped_count += 1
                    continue

                # Download newly added / modified file
                local_target = os.path.join(temp_receipts_dir, name)
                success, file_hash, byte_size = sp.download_file(rel_url, local_target)
                if success:
                    new_downloads_count += 1
                    folder_new_count += 1
                    new_bytes_transferred += byte_size

                    # Upload to Google Drive if connected
                    uploaded_id = None
                    if gdrive.is_connected and target_gdrive_folder_id:
                        uploaded_id = gdrive.upload_file(local_target, destination_folder_id=target_gdrive_folder_id, remote_file_name=name, overwrite=True)

                    if gdrive.is_connected and not uploaded_id:
                        print(f"⚠️ Failed to upload {name} to Google Drive; skipping manifest update so it will be retried next time.")
                    else:
                        # Update manifest
                        manifest_files[rel_url] = {
                            "name": name,
                            "size": byte_size,
                            "sha256": file_hash,
                            "category": folder_name,
                            "last_backed_up": today_str
                        }

                    # Remove temp file to conserve runner disk space
                    if os.path.exists(local_target):
                        os.remove(local_target)

                    # Periodic manifest checkpoint save every 50 files
                    if gdrive.is_connected and new_downloads_count % 50 == 0:
                        gdrive.save_manifest(manifest)

                total_in_folder = folder_new_count + folder_skipped_count
                if total_in_folder % 25 == 0 or total_in_folder == len(folder_files):
                    print(f"   ⏳ [{folder_name}] Progress: [{total_in_folder}/{len(folder_files)}] ({folder_new_count} new uploaded, {folder_skipped_count} skipped, {new_bytes_transferred / (1024*1024):.1f} MB total)...")

        # 4.2 Sync Income Folders (including bank and project nested subdirectories)
        if not (max_media_files and new_downloads_count >= max_media_files):
            print("\n⏳ Scanning SharePoint /sites/CHUNKING/DocLib/收入 folders...")

            # Helper to recursively discover subfolders with files under a given root
            def scan_income_folders(base_sp_path: str, current_rel_path: str = ""):
                folders_to_visit = [(base_sp_path, current_rel_path)]
                while folders_to_visit:
                    sp_path, rel_path = folders_to_visit.pop(0)
                    files = sp.fetch_folder_files_metadata(sp_path)
                    if files:
                        yield sp_path, rel_path, files
                    subdirs = sp.fetch_subfolders(sp_path)
                    for s in subdirs:
                        sub_name = os.path.basename(s.rstrip("/"))
                        next_rel = f"{rel_path}/{sub_name}".strip("/") if rel_path else sub_name
                        folders_to_visit.append((s, next_rel))

            gdrive_inc_folder_cache = {}

            def get_or_create_gdrive_path(rel_path: str, base_parent_id: str) -> str:
                if not rel_path:
                    return base_parent_id
                if rel_path in gdrive_inc_folder_cache:
                    return gdrive_inc_folder_cache[rel_path]
                parts = rel_path.split("/")
                curr = base_parent_id
                for part in parts:
                    curr = gdrive.find_or_create_folder(part, parent_id=curr)
                gdrive_inc_folder_cache[rel_path] = curr
                return curr

            for sp_folder_path, rel_path, sub_files in scan_income_folders("/sites/CHUNKING/DocLib/收入"):
                if max_media_files and new_downloads_count >= max_media_files:
                    break

                target_gdrive_sub_id = None
                existing_inc_drive = {}
                if gdrive.is_connected and gdrive_receipts_inc_id:
                    target_gdrive_sub_id = get_or_create_gdrive_path(rel_path, gdrive_receipts_inc_id)
                    existing_inc_drive = gdrive.list_files_in_folder(target_gdrive_sub_id)

                for file_meta in sub_files:
                    if max_media_files and new_downloads_count >= max_media_files:
                        break
                    name = file_meta.get("Name", "")
                    rel_url = file_meta.get("ServerRelativeUrl", "")
                    file_len = int(file_meta.get("Length", 0))

                    manifest_entry = manifest_files.get(rel_url)
                    if manifest_entry and manifest_entry.get("size") == file_len:
                        skipped_count += 1
                        continue

                    # Direct check against Google Drive
                    if name in existing_inc_drive and existing_inc_drive[name] == file_len:
                        manifest_files[rel_url] = {
                            "name": name,
                            "size": file_len,
                            "sha256": "pre-existing",
                            "category": f"Income/{rel_path}" if rel_path else "Income",
                            "last_backed_up": today_str
                        }
                        skipped_count += 1
                        continue

                    local_target = os.path.join(temp_receipts_dir, name)
                    success, file_hash, byte_size = sp.download_file(rel_url, local_target)
                    if success:
                        new_downloads_count += 1
                        new_bytes_transferred += byte_size
                        uploaded_id = None
                        if gdrive.is_connected and target_gdrive_sub_id:
                            uploaded_id = gdrive.upload_file(local_target, destination_folder_id=target_gdrive_sub_id, remote_file_name=name, overwrite=True)

                        if gdrive.is_connected and not uploaded_id:
                            print(f"⚠️ Failed to upload {name} to Google Drive; skipping manifest update so it will be retried next time.")
                        else:
                            manifest_files[rel_url] = {
                                "name": name,
                                "size": byte_size,
                                "sha256": file_hash,
                                "category": f"Income/{rel_path}" if rel_path else "Income",
                                "last_backed_up": today_str
                            }

                        if os.path.exists(local_target):
                            os.remove(local_target)

                        # Periodic manifest checkpoint save every 50 files
                        if gdrive.is_connected and new_downloads_count % 50 == 0:
                            gdrive.save_manifest(manifest)

                    if (new_downloads_count + skipped_count) % 25 == 0:
                        disp_name = f"Income/{rel_path}" if rel_path else "Income"
                        print(f"   ⏳ [{disp_name}] Progress: ({new_downloads_count} new uploaded, {skipped_count} skipped, {new_bytes_transferred / (1024*1024):.1f} MB total)...")

    print(f"✅ Incremental Media Sync Completed:")
    print(f"   ⏩ Skipped (already backed up): {skipped_count} files")
    print(f"   📥 Newly synced files: {new_downloads_count} files ({new_bytes_transferred / (1024*1024):.2f} MB)")

    # Save manifest
    manifest["last_sync_date"] = today_str
    manifest["total_tracked_files"] = len(manifest_files)
    if gdrive.is_connected:
        gdrive.save_manifest(manifest)
    else:
        with open(os.path.join(SCRIPT_DIR, "manifest.json"), "w", encoding="utf-8") as f:
            json.dump(manifest, f, ensure_ascii=False, indent=2)

    # Clean up empty temp download directory
    try:
        os.rmdir(temp_receipts_dir)
    except OSError:
        pass

    # =========================================================================
    # 5. AUDIT REPORT & SUMMARY
    # =========================================================================
    elapsed = time.time() - start_time
    report = {
        "status": "SUCCESS",
        "timestamp": timestamp_full,
        "duration_seconds": round(elapsed, 2),
        "records": {
            "projects": len(clean_projects),
            "expenses": len(clean_expenses),
            "incomes": len(clean_incomes)
        },
        "media_sync": {
            "new_files_transferred": new_downloads_count,
            "new_megabytes_transferred": round(new_bytes_transferred / (1024*1024), 2),
            "skipped_already_synced": skipped_count,
            "total_tracked_receipts": len(manifest_files)
        },
        "throttling_metrics": {
            "total_429_or_503_events": sp.total_throttle_events,
            "circuit_breaker_tripped": False
        }
    }

    report_path = os.path.join(staging_dir, "backup_report.json")
    with open(report_path, "w", encoding="utf-8") as f:
        json.dump(report, f, ensure_ascii=False, indent=2)

    # Also save to root of backup dir for GitHub Actions artifact upload
    with open(os.path.join(SCRIPT_DIR, "backup_report.json"), "w", encoding="utf-8") as f:
        json.dump(report, f, ensure_ascii=False, indent=2)

    if gdrive.is_connected:
        gdrive.upload_file(report_path, destination_folder_id=gdrive_weekly_folder_id, overwrite=True)

    print("\n==================================================================")
    print(f"🎉 BACKUP COMPLETED IN {elapsed:.2f}s WITH STATUS: {report['status']}")
    print(f"📊 Projects: {len(clean_projects)} | Expenses: {len(clean_expenses)} | Income: {len(clean_incomes)}")
    print(f"💾 New Media: {new_downloads_count} files ({new_bytes_transferred / (1024*1024):.2f} MB)")
    print(f"🛡️ Throttling Events (HTTP 429/503): {sp.total_throttle_events} (Paced & Handled smoothly)")
    print("==================================================================")


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(description="CHUN KING SharePoint to Google Drive Backup Pipeline")
    parser.add_argument("--lists-only", action="store_true", help="Only backup SharePoint lists (Projects, Expenses, Income) and skip media files")
    parser.add_argument("--max-media-files", type=int, default=None, help="Maximum number of media files to sync (for testing)")
    args = parser.parse_args()

    try:
        run_backup_pipeline(lists_only=args.lists_only, max_media_files=args.max_media_files)
    except SharePointThrottlingException as te:
        print(f"\n❌ FATAL: SharePoint throttling circuit breaker triggered:\n{te}")
        sys.exit(2)
    except Exception as e:
        print(f"\n❌ FATAL BACKUP FAILURE: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)
